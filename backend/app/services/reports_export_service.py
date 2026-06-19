from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.reports_service import ReportsService


class ReportsExportService:

    @staticmethod
    def create_asset_status_excel():

        report = ReportsService.get_asset_status_report()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Asset Status Report"

        headers = [
            "Status",
            "Count"
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9EAF7",
                end_color="D9EAF7",
                fill_type="solid"
            )

        for row in report:
            sheet.append([
                row["status"],
                row["count"]
            ])

        sheet.column_dimensions["A"].width = 25
        sheet.column_dimensions["B"].width = 15

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return output
    
    @staticmethod
    def create_asset_category_excel():

        report = ReportsService.get_asset_category_report()

        workbook = Workbook()
        sheet = workbook.active

        sheet.title = "Asset Category Report"

        headers = [
            "Category",
            "Count"
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9EAF7",
                end_color="D9EAF7",
                fill_type="solid"
            )

        for row in report:
            sheet.append([
                row["category"],
                row["count"]
            ])

        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 15

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return output
    
    @staticmethod
    def create_employee_asset_excel():

        report = ReportsService.get_employee_asset_report()

        workbook = Workbook()
        sheet = workbook.active

        sheet.title = "Employee Asset Report"

        headers = [
            "Employee Code",
            "Employee Name",
            "Asset Count"
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9EAF7",
                end_color="D9EAF7",
                fill_type="solid"
            )

        for row in report:
            sheet.append([
                row["employee_code"],
                row["employee"],
                row["asset_count"]
            ])

        sheet.column_dimensions["A"].width = 20
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 15

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return output
    
    @staticmethod
    def create_all_assets_excel():

        from app.models.asset_model import Asset

        assets = Asset.query.filter_by(
            is_active=True
        ).all()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "All Assets"

        headers = [
            "Asset Code",
            "Asset Name",
            "Serial Number",
            "Category",
            "Vendor",
            "Location",
            "Employee",
            "Status",
            "Purchase Cost",
            "Invoice Number"
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                start_color="D9EAF7",
                end_color="D9EAF7",
                fill_type="solid"
            )

        for asset in assets:
            sheet.append([
                asset.asset_code,
                asset.asset_name,
                asset.serial_number,
                asset.category.name if asset.category else None,
                asset.vendor.name if asset.vendor else None,
                asset.location.name if asset.location else None,
                asset.employee.full_name if asset.employee else None,
                asset.status,
                asset.purchase_cost,
                asset.invoice_number
            ])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return output